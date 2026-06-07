"""Document loaders (KB-05).

Turn uploaded file bytes into plain text + a source-type hint, dispatched by
extension. Markdown / text / source code are read directly; PDF and DOCX use
PyMuPDF / python-docx (imported lazily so the module loads without them).

Unsupported types raise ``UnsupportedFileTypeError`` — the upload UI rejects
them client-side too (frozen mock Screen 01), this is the server backstop.
"""

from __future__ import annotations

from pathlib import PurePosixPath

# Extensions → source_type for the chunker's kind detection.
_TEXT_EXT = {".md", ".markdown", ".txt", ".rst", ".csv", ".tsv"}
_CODE_EXT = {
    ".py", ".ts", ".tsx", ".js", ".jsx", ".go", ".java", ".rs", ".c", ".cc",
    ".cpp", ".h", ".hpp", ".rb", ".php", ".cs", ".kt", ".swift", ".sql",
    ".sh", ".yaml", ".yml", ".toml", ".json",
}
_PDF_EXT = {".pdf"}
_DOCX_EXT = {".docx"}
_XLSX_EXT = {".xlsx"}

SUPPORTED_EXTENSIONS = _TEXT_EXT | _CODE_EXT | _PDF_EXT | _DOCX_EXT | _XLSX_EXT

# Spreadsheets can be huge — cap what we flatten into text so a stray 100k-row
# export doesn't blow up ingest. Generous but bounded.
_XLSX_MAX_ROWS_PER_SHEET = 2000


class UnsupportedFileTypeError(ValueError):
    """Raised for a file extension the loaders can't handle."""


class LoaderUnavailableError(RuntimeError):
    """Raised when the optional lib for a supported type isn't installed."""


def _ext(filename: str) -> str:
    return PurePosixPath(filename).suffix.lower()


def source_type_for(filename: str) -> str:
    """Classify a filename into a KB ``source_type`` used downstream by the
    chunker's kind detection."""
    ext = _ext(filename)
    if ext in _CODE_EXT:
        return "code"
    return "upload"


def load_text(filename: str, data: bytes) -> tuple[str, str]:
    """Return ``(text, source_type)`` for an uploaded file.

    Raises ``UnsupportedFileTypeError`` for unknown extensions and
    ``LoaderUnavailableError`` if a needed optional lib is missing.
    """
    ext = _ext(filename)
    if ext in _TEXT_EXT or ext in _CODE_EXT:
        return _decode(data), source_type_for(filename)
    if ext in _PDF_EXT:
        return _load_pdf(data), "upload"
    if ext in _DOCX_EXT:
        return _load_docx(data), "upload"
    if ext in _XLSX_EXT:
        return _load_xlsx(data), "upload"
    raise UnsupportedFileTypeError(
        f"unsupported file type '{ext or filename}'. Supported: "
        f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    )


def _decode(data: bytes) -> str:
    """Best-effort UTF-8 decode with a latin-1 fallback (never raises)."""
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="replace")


def _load_pdf(data: bytes) -> str:
    try:
        import fitz  # type: ignore[import-untyped]  # PyMuPDF
    except Exception as e:  # noqa: BLE001
        raise LoaderUnavailableError(f"PDF support needs PyMuPDF: {e}") from e
    parts: list[str] = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for page in doc:
            parts.append(page.get_text())
    return "\n\n".join(parts).strip()


def _load_docx(data: bytes) -> str:
    try:
        import io

        import docx  # python-docx
    except Exception as e:  # noqa: BLE001
        raise LoaderUnavailableError(f"DOCX support needs python-docx: {e}") from e
    document = docx.Document(io.BytesIO(data))
    return "\n\n".join(p.text for p in document.paragraphs if p.text.strip()).strip()


def _load_xlsx(data: bytes) -> str:
    """Flatten a workbook to text: one ``## Sheet: <name>`` heading per sheet,
    then each non-empty row as a tab-joined line. Read-only + bounded rows so a
    huge export can't blow up ingest. Empty cells are kept as blanks to preserve
    column alignment within a row."""
    try:
        import io

        import openpyxl  # optional dep
    except Exception as e:  # noqa: BLE001
        raise LoaderUnavailableError(f"XLSX support needs openpyxl: {e}") from e
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for ws in wb.worksheets:
            parts.append(f"## Sheet: {ws.title}")
            rows_out = 0
            for row in ws.iter_rows(values_only=True):
                if rows_out >= _XLSX_MAX_ROWS_PER_SHEET:
                    parts.append(f"… (truncated at {_XLSX_MAX_ROWS_PER_SHEET} rows)")
                    break
                cells = ["" if c is None else str(c) for c in row]
                if any(c.strip() for c in cells):  # skip fully-blank rows
                    parts.append("\t".join(cells).rstrip())
                    rows_out += 1
            parts.append("")
    finally:
        wb.close()
    return "\n".join(parts).strip()
